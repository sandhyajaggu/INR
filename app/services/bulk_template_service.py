"""Shared builder for every bulk-upload template (.xlsx).

Columns are derived from the same Pydantic schema the matching bulk-upload
endpoint validates rows against, so a template can't drift from its upload
rules. Headers are chosen so the upload's header normalization
("EPIC No" -> epic_no) maps them straight back to the schema field names.

Every header carries a note (hover in Excel) with its rule. Single-sheet
templates also get example rows and an "Instructions" sheet after the data
sheet (uploads read the first/active sheet only). The combined beneficiaries
workbook gets neither: that upload treats every non-empty tab as a scheme,
so tabs stay header-only and unused tabs are skipped.
"""

import io
import re
import types
import typing
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.geography import Mandal
from app.services.excel_import_service import MAX_ROWS_PER_SHEET

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMATTED_ROWS = 5000  # rows pre-formatted as text/date and given dropdowns

HEADER_FONT = Font(bold=True, color="FFFFFF")
REQUIRED_FILL = PatternFill("solid", fgColor="C53030")
OPTIONAL_FILL = PatternFill("solid", fgColor="0B1F3A")

HEADER_OVERRIDES = {
    "epic_no": "EPIC No",
    "ifsc_code": "IFSC Code",
    "photo_url": "Photo URL",
    "video_url": "Video URL",
    "document_url": "Document URL",
}

# Identifier-like columns: stored as text so Excel keeps leading zeros and
# doesn't turn long numbers (Aadhaar, account numbers) into 1.23E+11.
TEXT_KEYS = {
    "epic_no", "aadhaar_number", "mobile", "mobile_number", "bank_account_number", "ifsc_code",
    "survey_number", "ration_card_number", "gas_connection_number", "bus_pass_number",
    "house_no", "booth_number", "booth_no",
}

RULES = {
    "epic_no": "3 letters + 7 digits, e.g. ABC1234567",
    "aadhaar_number": "Exactly 12 digits (stored encrypted)",
    "mobile": "10 digits, starts with 6-9",
    "mobile_number": "10 digits, starts with 6-9",
    "bank_account_number": "Bank account number",
    "ifsc_code": "Bank IFSC code, e.g. SBIN0001234",
    "amount": "Amount in rupees, more than 0",
    "estimated_cost": "Amount in rupees",
    "mandal_name": "One of the mandals in the dropdown",
    "village_name": "Village in that mandal",
    "photo_url": "Optional here: leave blank and add the photo later by editing the record",
    "video_url": "Optional here: leave blank and add the video later by editing the record",
    "document_url": "Link of the document, from the File Upload (POST /files/upload)",
    "land_extent_acres": "Acres, e.g. 2.5",
}

SAMPLE_VALUES: dict[str, object] = {
    "beneficiary_name": "Ravi Kumar",
    "farmer_name": "Ravi Kumar",
    "head_of_household_name": "Ravi Kumar",
    "mother_name": "Lakshmi Devi",
    "relation_name": "Venkat Rao",
    "student_name": "Sai Kumar",
    "school_name": "ZPHS Kandukur",
    "class_grade": "8",
    "qualification": "B.Tech",
    "epic_no": "ABC1234567",
    "age": 42,
    "aadhaar_number": "123412341234",
    "mobile_number": "9876543210",
    "bank_account_number": "30012345678",
    "ifsc_code": "SBIN0001234",
    "amount": 25000,
    "mandal_name": "Kandukur",
    "village_name": "Palukur",
    "application_date": date(2026, 9, 1),
    "status": "approved",
    "ration_card_number": "WAP123456789",
    "gas_connection_number": "GC1002345",
    "gas_agency": "Bharat Gas, Kandukur",
    "bus_pass_number": "BP2026001",
    "preferred_route": "Kandukur - Ongole",
    "depot": "Kandukur",
    "land_extent_acres": 2.5,
    "survey_number": "123/4A",
    # development works
    "title": "CC road, Palukur SC colony",
    "category": "Roads",
    "estimated_cost": 1500000,
    "description": "400 m cement concrete road",
    "work_date": date(2026, 9, 1),
}


@dataclass
class TemplateColumn:
    header: str
    required: bool
    rule: str
    as_text: bool = False
    as_date: bool = False
    choices: list[str] | None = None


@dataclass
class TemplateSheet:
    title: str
    columns: list[TemplateColumn]
    example_rows: list[list[object]] = field(default_factory=list)


def header_for(key: str) -> str:
    return HEADER_OVERRIDES.get(key, key.replace("_", " ").title())


def _base_type(annotation: object) -> object:
    if isinstance(annotation, types.UnionType) or typing.get_origin(annotation) is typing.Union:
        args = [a for a in typing.get_args(annotation) if a is not type(None)]
        return args[0] if args else str
    return annotation


def _pattern_choices(metadata: list) -> list[str] | None:
    for item in metadata:
        pattern = getattr(item, "pattern", None)
        match = re.fullmatch(r"\^\(([^()]+)\)\$", pattern or "")
        if match:
            return match.group(1).split("|")
    return None


def columns_from_schema(schema: type[BaseModel], mandals: list[str]) -> list[TemplateColumn]:
    columns = []
    for key, info in schema.model_fields.items():
        base = _base_type(info.annotation)
        choices = _pattern_choices(info.metadata)
        if key == "mandal_name" and mandals:
            choices = mandals
        rule = RULES.get(key)
        if rule is None:
            if choices:
                rule = " / ".join(choices)
            elif base is date:
                rule = "Date: YYYY-MM-DD or an Excel date"
            elif base is int:
                rule = "Whole number"
            elif base in (float, Decimal):
                rule = "Number"
            else:
                rule = "Text"
        if choices and key != "mandal_name":
            rule = f"One of: {', '.join(choices)}"
            if not info.is_required() and info.default is not None:
                rule += f" (blank = {info.default})"
        columns.append(
            TemplateColumn(
                header=header_for(key),
                required=info.is_required(),
                rule=rule,
                as_text=key in TEXT_KEYS,
                as_date=base is date,
                choices=choices,
            )
        )
    return columns


def example_row(schema: type[BaseModel]) -> list[object]:
    row = []
    for key, info in schema.model_fields.items():
        value = SAMPLE_VALUES.get(key)
        choices = _pattern_choices(info.metadata)
        # A shared sample (e.g. status "approved") may not be valid for this schema.
        if choices and value not in choices:
            value = choices[0]
        row.append(value)
    return row


async def mandal_names(db: AsyncSession) -> list[str]:
    return list((await db.execute(select(Mandal.name).order_by(Mandal.name))).scalars().all())


def _fill_sheet(sheet, spec: TemplateSheet) -> None:
    sheet.append([c.header for c in spec.columns])
    for row in spec.example_rows:
        sheet.append(row)
    for index, column in enumerate(spec.columns, start=1):
        cell = sheet.cell(1, index)
        cell.font = HEADER_FONT
        cell.fill = REQUIRED_FILL if column.required else OPTIONAL_FILL
        cell.comment = Comment(f"{'Required' if column.required else 'Optional'}. {column.rule}", "INR MLA CRM")
        letter = cell.column_letter
        sheet.column_dimensions[letter].width = max(14, len(column.header) + 4)
        if column.as_text or column.as_date:
            fmt = "@" if column.as_text else "yyyy-mm-dd"
            for r in range(2, FORMATTED_ROWS + 2):
                sheet.cell(r, index).number_format = fmt
        if column.choices:
            validation = DataValidation(type="list", formula1='"' + ",".join(column.choices) + '"', allow_blank=True)
            validation.add(f"{letter}2:{letter}{FORMATTED_ROWS + 1}")
            sheet.add_data_validation(validation)
    sheet.freeze_panes = "A2"


def _add_instructions(workbook: Workbook, spec: TemplateSheet, notes: list[str]) -> None:
    guide = workbook.create_sheet("Instructions")
    guide.append(["Column", "Required", "Rule"])
    for cell in guide[1]:
        cell.font = HEADER_FONT
        cell.fill = OPTIONAL_FILL
    for column in spec.columns:
        guide.append([column.header, "Yes" if column.required else "No", column.rule])
    guide.append([])
    guide.append(["Notes", "", " ".join(notes)])
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 90
    for row in guide.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")


def standard_notes(extra: list[str] | None = None) -> list[str]:
    return [
        "Red headers are required; hover a header to see its rule.",
        "Delete the example row(s) before uploading.",
        f"Max {MAX_ROWS_PER_SHEET:,} rows per file; split larger lists by mandal.",
        "If any row has a problem, nothing is saved and every problem is listed with its row number.",
        *(extra or []),
    ]


def build_single_sheet_template(spec: TemplateSheet, notes: list[str]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = spec.title
    _fill_sheet(sheet, spec)
    _add_instructions(workbook, spec, notes)
    workbook.active = 0
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_multi_sheet_template(specs: list[TemplateSheet]) -> bytes:
    """One tab per spec, headers only (see module docstring)."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for spec in specs:
        _fill_sheet(workbook.create_sheet(spec.title), spec)
    workbook.active = 0
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def xlsx_attachment_headers(filename: str) -> dict[str, str]:
    return {"Content-Disposition": f'attachment; filename="{filename}"'}

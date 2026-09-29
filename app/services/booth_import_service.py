"""Booths bulk upload: sheet parsing (.xlsx / .xls / .csv), validation, upsert.

All-or-nothing like the voters import: every row is validated first
(required fields, numbers, status, resolvable mandal/village, no duplicate
booth within the file). If any row fails, nothing is written. Otherwise a
row whose (mandal, booth number) already exists updates that booth, and
every other row adds a new booth.
"""

import csv
import io

import xlrd
from fastapi import HTTPException, UploadFile, status
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.geography import Booth
from app.schemas.booth_bulk import BOOTH_COLUMN_LABELS, BOOTH_SHEET_COLUMNS, BoothBulkRow
from app.schemas.bulk_import import BulkImportResult, BulkImportRowError
from app.services.activity_service import log_activity
from app.services.booth_service import _booth_rows_query, _ordered
from app.services.excel_import_service import (
    MAX_ROWS_PER_SHEET,
    _load_workbook_from_upload,
    _normalize_header,
    _rows_from_iter,
)
from app.services.geography_service import load_geography_maps

TEMPLATE_EXAMPLE_ROWS = [
    ["001", "Ravi Kumar", "Kandukur", "Palur", 1000, 420, 210, 150, 30, 30, "On Time"],
    ["002", "Suresh Babu", "Kandukur", "Ogur", 950, 380, 160, 180, 10, 30, "Delayed"],
]


def _unprocessable(message: str) -> HTTPException:
    return HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, message)


def _blank_to_none(value: object) -> object:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _check_required_columns(headers: set[str]) -> None:
    missing = [label for label, key in BOOTH_SHEET_COLUMNS.items() if key not in headers]
    if missing:
        raise _unprocessable(f"Missing required column(s): {', '.join(missing)}")


def _rows_from_table(table: list[list[object]]) -> list[dict[str, object]]:
    """Header check + row dicts (same shape as parse_excel_rows) for every upload format."""
    if not table:
        raise _unprocessable("The sheet is empty")
    header_row, *data_rows = table
    headers = [_normalize_header(h) for h in header_row if h is not None]
    _check_required_columns(set(headers))
    rows = _rows_from_iter(iter(data_rows), headers)
    if not rows:
        raise _unprocessable("The sheet has no booth rows below the header")
    if len(rows) > MAX_ROWS_PER_SHEET:
        raise _unprocessable(
            f"Sheet has {len(rows)} rows, which exceeds the {MAX_ROWS_PER_SHEET:,}-row limit per "
            "upload. Split it into smaller files and upload each one separately."
        )
    return rows


def _read_csv(contents: bytes) -> list[list[object]]:
    try:
        text = contents.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel's "Save as CSV" on Windows writes the ANSI code page, not UTF-8.
        text = contents.decode("cp1252", errors="replace")
    return [[_blank_to_none(v) for v in row] for row in csv.reader(io.StringIO(text))]


def _read_xlsx(contents: bytes) -> list[list[object]]:
    sheet = _load_workbook_from_upload(contents).active
    return [[_blank_to_none(v) for v in row] for row in sheet.iter_rows(values_only=True)]


def _read_xls(contents: bytes) -> list[list[object]]:
    try:
        book = xlrd.open_workbook(file_contents=contents)
    except Exception as exc:
        raise _unprocessable("Could not read the Excel file") from exc
    sheet = book.sheet_by_index(0)
    table = []
    for r in range(sheet.nrows):
        row = []
        for value in sheet.row_values(r):
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            row.append(_blank_to_none(value))
        table.append(row)
    return table


async def parse_booth_upload(file: UploadFile) -> list[dict[str, object]]:
    name = (file.filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _rows_from_table(_read_xlsx(await file.read()))
    if name.endswith(".xls"):
        return _rows_from_table(_read_xls(await file.read()))
    if name.endswith(".csv"):
        return _rows_from_table(_read_csv(await file.read()))
    raise _unprocessable("Only .xlsx, .xls and .csv files are supported")


def _format_validation_error(exc: ValidationError) -> str:
    messages = []
    for err in exc.errors():
        label = BOOTH_COLUMN_LABELS.get(str(err["loc"][0]), "") if err["loc"] else ""
        msg = err["msg"].removeprefix("Value error, ")
        if err["type"] == "missing":
            msg = "is required"
        elif err["type"].startswith("int_"):
            msg = f"must be a whole number, got '{err['input']}'"
        messages.append(f"{label} {msg}".strip())
    return "; ".join(messages)


async def bulk_import_booths(db: AsyncSession, rows: list[dict], actor_id: int) -> BulkImportResult:
    errors: list[BulkImportRowError] = []
    parsed_rows: list[tuple[int, BoothBulkRow]] = []

    for raw in rows:
        row_num = raw.get("_row_number")
        try:
            parsed_rows.append((row_num, BoothBulkRow.model_validate(raw)))
        except ValidationError as exc:
            booth_hint = raw.get("booth_no")
            prefix = f"Booth No '{booth_hint}': " if booth_hint is not None else ""
            errors.append(BulkImportRowError(row=row_num, reason=prefix + _format_validation_error(exc)))

    mandal_map, village_map, _ = await load_geography_maps(db)
    resolved: list[tuple[BoothBulkRow, int, int]] = []
    first_seen: dict[tuple[int, str], int] = {}
    for row_num, parsed in parsed_rows:
        prefix = f"Booth No '{parsed.booth_no}': "
        mandal_id = mandal_map.get(parsed.mandal.lower())
        if mandal_id is None:
            errors.append(BulkImportRowError(row=row_num, reason=f"{prefix}Unknown mandal '{parsed.mandal}'"))
            continue
        village_id = village_map.get((mandal_id, parsed.village.lower()))
        if village_id is None:
            errors.append(
                BulkImportRowError(
                    row=row_num, reason=f"{prefix}Unknown village '{parsed.village}' in mandal '{parsed.mandal}'"
                )
            )
            continue
        key = (mandal_id, parsed.booth_no)
        if key in first_seen:
            errors.append(
                BulkImportRowError(
                    row=row_num,
                    reason=f"{prefix}Duplicate booth in mandal '{parsed.mandal}' (first seen on row {first_seen[key]})",
                )
            )
            continue
        first_seen[key] = row_num
        resolved.append((parsed, mandal_id, village_id))

    if errors:
        errors.sort(key=lambda e: e.row)
        return BulkImportResult(inserted=0, errors=errors)

    # The booths table is small (a few hundred rows per constituency), so one
    # query for all of them beats a lookup per sheet row.
    existing = {
        (b.mandal_id, b.booth_number.strip()): b for b in (await db.execute(select(Booth))).scalars().all()
    }
    new_booths = []
    updated_count = 0
    for parsed, mandal_id, village_id in resolved:
        values = {
            "village_id": village_id,
            "booth_officer_name": parsed.in_charge,
            "total_voters": parsed.registered_votes,
            "votes_polled": parsed.votes_polled,
            "tdp_votes": parsed.tdp,
            "ysp_votes": parsed.ysp,
            "janasena_votes": parsed.janasena,
            "congress_votes": parsed.congress,
            "status": parsed.status,
        }
        booth = existing.get((mandal_id, parsed.booth_no))
        if booth is not None:
            for field, value in values.items():
                setattr(booth, field, value)
            updated_count += 1
        else:
            new_booths.append(Booth(booth_number=parsed.booth_no, mandal_id=mandal_id, **values))

    db.add_all(new_booths)
    await db.flush()
    await log_activity(
        db,
        actor_id=actor_id,
        action_type="booth_bulk_imported",
        module="booths",
        reference_id=None,
        description=f"Bulk import: {len(new_booths)} booths added, {updated_count} updated",
    )
    await db.commit()
    return BulkImportResult(inserted=len(new_booths), updated=updated_count, errors=[])


async def build_booth_template(db: AsyncSession) -> bytes:
    """The upload template, pre-filled with every current booth.

    Same 11 columns the upload accepts, so the file can be edited and
    uploaded straight back (matching Mandal + Booth No rows update). Falls
    back to example rows when there are no booths yet. A booth added
    without an In-Charge or Status exports those cells blank; the upload
    asks for them before it accepts the row.
    """
    booths = [
        [b.booth_number, b.booth_officer_name, mandal, village, b.total_voters, b.votes_polled,
         b.tdp_votes, b.ysp_votes, b.janasena_votes, b.congress_votes, b.status]
        for b, mandal, village in (await db.execute(_ordered(_booth_rows_query()))).all()
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Booths"
    sheet.append(list(BOOTH_SHEET_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B1F3A")
    for row in booths or TEMPLATE_EXAMPLE_ROWS:
        sheet.append(row)
    # Text format on Booth No keeps leading zeros ("001") when users type new rows.
    for r in range(2, sheet.max_row + 1000):
        sheet.cell(r, 1).number_format = "@"
    for col, width in zip("ABCDEFGHIJK", (10, 22, 18, 20, 17, 13, 8, 8, 10, 10, 10)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

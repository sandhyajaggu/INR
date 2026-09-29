"""Voters bulk-upload template (.xlsx) for POST /voters/bulk-upload.

A blank template with example rows, not a pre-filled export: the voter roll
(~230k rows) is far larger than one download or one upload (capped at
MAX_ROWS_PER_SHEET) can hold. Headers normalize to exactly the keys
VoterBulkRow expects ("EPIC No" -> epic_no, ...). The "Voters" sheet is the
active one, which is the sheet the upload reads; "Instructions" follows it.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.geography import Mandal
from app.services.excel_import_service import MAX_ROWS_PER_SHEET

# (header, required, stored as text, rule shown on the Instructions sheet)
VOTER_TEMPLATE_COLUMNS = [
    ("EPIC No", True, True, "3 letters + 7 digits, e.g. ABC1234567. An existing EPIC No updates that voter."),
    ("Name", True, False, "Voter's full name"),
    ("Relation Name", False, False, "Father / husband name"),
    ("Age", False, False, "Whole number, 0-130"),
    ("Gender", False, False, "Male, Female or Other"),
    ("Mobile", False, True, "10 digits, starts with 6-9"),
    ("Aadhaar Number", False, True, "Exactly 12 digits (stored encrypted)"),
    ("House No", False, True, "As printed on the roll"),
    ("Mandal Name", True, False, "One of the mandals in the dropdown"),
    ("Village Name", True, False, "Village in that mandal"),
    ("Booth Number", False, True, "Must already exist in that mandal (Booths page)"),
    ("Voted Last Election", False, False, "Yes or No"),
    ("Is New Voter", False, False, "Yes or No (blank = No)"),
]
EXAMPLE_ROWS = [
    ["ABC1234567", "Ravi Kumar", "Venkat Rao", 42, "Male", "9876543210", "123412341234", "4-12",
     "Kandukur", "Palukur", "81", "Yes", "No"],
    ["XYZ7654321", "Lakshmi Devi", "Srinivas", 19, "Female", None, None, "7-3A",
     "Lingasamudram", "Chinapavani", "59", None, "Yes"],
]
FORMATTED_ROWS = 5000  # rows pre-formatted as text / given dropdowns

HEADER_FONT = Font(bold=True, color="FFFFFF")
REQUIRED_FILL = PatternFill("solid", fgColor="C53030")
OPTIONAL_FILL = PatternFill("solid", fgColor="0B1F3A")


def _add_list_validation(sheet, values: list[str], column_letter: str) -> None:
    validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    validation.add(f"{column_letter}2:{column_letter}{FORMATTED_ROWS + 1}")
    sheet.add_data_validation(validation)


async def build_voter_template(db: AsyncSession) -> bytes:
    mandals = (await db.execute(select(Mandal.name).order_by(Mandal.name))).scalars().all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Voters"
    sheet.append([header for header, *_ in VOTER_TEMPLATE_COLUMNS])
    for row in EXAMPLE_ROWS:
        sheet.append(row)

    letters = {}
    for index, (header, required, as_text, _) in enumerate(VOTER_TEMPLATE_COLUMNS, start=1):
        cell = sheet.cell(1, index)
        cell.font = HEADER_FONT
        cell.fill = REQUIRED_FILL if required else OPTIONAL_FILL
        letters[header] = cell.column_letter
        sheet.column_dimensions[cell.column_letter].width = max(14, len(header) + 4)
        # Text format stops Excel turning "0081" into 81 or a 12-digit Aadhaar into 1.23E+11.
        if as_text:
            for r in range(2, FORMATTED_ROWS + 2):
                sheet.cell(r, index).number_format = "@"

    _add_list_validation(sheet, ["Male", "Female", "Other"], letters["Gender"])
    if mandals:
        _add_list_validation(sheet, list(mandals), letters["Mandal Name"])
    _add_list_validation(sheet, ["Yes", "No"], letters["Voted Last Election"])
    _add_list_validation(sheet, ["Yes", "No"], letters["Is New Voter"])
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet("Instructions")
    guide.append(["Column", "Required", "Rule"])
    for cell in guide[1]:
        cell.font = HEADER_FONT
        cell.fill = OPTIONAL_FILL
    for header, required, _, rule in VOTER_TEMPLATE_COLUMNS:
        guide.append([header, "Yes" if required else "No", rule])
    guide.append([])
    guide.append([
        "Notes", "",
        "Red headers are required. Delete the 2 example rows before uploading. "
        f"Max {MAX_ROWS_PER_SHEET:,} rows per file; split larger lists by mandal. "
        "If any row has a problem, nothing is saved and every problem is listed with its row number.",
    ])
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 90
    for row in guide.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    workbook.active = 0
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
